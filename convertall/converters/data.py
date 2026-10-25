"""Structured data: JSON, JSON Lines, YAML, TOML, XML, INI, plist, and tables (CSV, TSV, XLSX, SQLite)."""

import configparser
import csv
import datetime as dt
import json
import plistlib
import re
import sqlite3
import tomllib
import xml.etree.ElementTree as ET
from pathlib import Path

import tomli_w
import yaml

from ..errors import ConversionError
from ..registry import converter

TREE = {"json", "jsonl", "yaml", "toml", "xml", "ini", "plist"}
TABLE = {"csv", "tsv", "xlsx"}
READ_ONLY = {"sqlite"}


# ---- XML <-> dict (attributes as "@name", text as "#text") ----

def _xml_to_obj(el: ET.Element):
    obj: dict = {f"@{k}": v for k, v in el.attrib.items()}
    for child in el:
        val = _xml_to_obj(child)
        if child.tag in obj:
            if not isinstance(obj[child.tag], list):
                obj[child.tag] = [obj[child.tag]]
            obj[child.tag].append(val)
        else:
            obj[child.tag] = val
    text = (el.text or "").strip()
    if not obj:
        return text or None
    if text:
        obj["#text"] = text
    return obj


def _obj_to_xml(parent: ET.Element, tag: str, val) -> None:
    if isinstance(val, list):
        for item in val:
            _obj_to_xml(parent, tag, item)
        return
    el = ET.SubElement(parent, _safe_tag(tag))
    if isinstance(val, dict):
        for k, v in val.items():
            if k.startswith("@"):
                el.set(k[1:], _scalar(v))
            elif k == "#text":
                el.text = _scalar(v)
            else:
                _obj_to_xml(el, k, v)
    elif val is not None:
        el.text = _scalar(val)


def _safe_tag(tag) -> str:
    tag = "".join(c if c.isalnum() or c in "-_." else "_" for c in str(tag))
    return tag if tag and (tag[0].isalpha() or tag[0] == "_") else f"_{tag}"


def _scalar(v) -> str:
    if isinstance(v, bool):
        return str(v).lower()
    return "" if v is None else str(v)


# ---- load / dump ----

def load(path: Path, fmt: str):
    text = lambda: path.read_text(encoding="utf-8-sig")  # noqa: E731
    if fmt == "json":
        return json.loads(text())
    if fmt == "jsonl":
        return [json.loads(line) for line in text().splitlines() if line.strip()]
    if fmt == "plist":
        with path.open("rb") as f:
            return plistlib.load(f)  # XML or binary
    if fmt == "sqlite":
        return _load_sqlite(path)
    if fmt == "yaml":
        return yaml.safe_load(text())
    if fmt == "toml":
        return tomllib.loads(text())
    if fmt == "xml":
        root = ET.fromstring(text())
        return {root.tag: _xml_to_obj(root)}
    if fmt == "ini":
        # no default-section magic: [DEFAULT] is kept as written instead of copied into every section
        cp = configparser.ConfigParser(interpolation=None, default_section="\0", strict=False)
        cp.optionxform = str  # keep key case
        try:
            cp.read_string(text())
        except configparser.MissingSectionHeaderError:  # keys before any [section]
            cp.read_string("[DEFAULT]\n" + text())
        return {s: dict(cp[s]) for s in cp.sections()}
    if fmt in ("csv", "tsv"):
        with path.open(newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096)
            f.seek(0)
            delim = "\t" if fmt == "tsv" else _sniff_delim(sample)
            return list(csv.DictReader(f, delimiter=delim))
    if fmt == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = {ws.title: _rows_to_records(list(ws.iter_rows(values_only=True))) for ws in wb.worksheets}
        wb.close()
        return next(iter(sheets.values())) if len(sheets) == 1 else sheets
    raise ConversionError(f"Can't read {fmt}")


def _load_sqlite(path: Path):
    con = sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)
    try:
        names = [r[0] for r in con.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY rowid")]
        tables = {}
        for name in names:
            cur = con.execute('SELECT * FROM "{}"'.format(name.replace('"', '""')))
            cols = [d[0] for d in cur.description]
            tables[name] = [{c: (f"<{len(v)} bytes>" if isinstance(v, bytes) else v) for c, v in zip(cols, row)}
                            for row in cur]
    finally:
        con.close()
    if not tables:
        raise ConversionError("This database has no tables")
    return next(iter(tables.values())) if len(tables) == 1 else tables


_NUMBER = re.compile(r"-?(?:0|[1-9]\d{0,14})(?:\.\d+)?")


def _numbers(data):
    """CSV cells are all text; give Excel real numbers ('007' and long IDs stay text)."""
    def conv(v):
        if isinstance(v, str) and _NUMBER.fullmatch(v):
            return float(v) if "." in v else int(v)
        return v
    if isinstance(data, list):
        return [{k: conv(v) for k, v in r.items()} if isinstance(r, dict) else r for r in data]
    return data


def _sniff_delim(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _rows_to_records(rows: list[tuple]) -> list[dict]:
    rows = [r for r in rows if any(c is not None for c in r)]
    if not rows:
        return []
    header = [str(h) if h is not None else f"column{i + 1}" for i, h in enumerate(rows[0])]
    return [{h: _plain(v) for h, v in zip(header, r)} for r in rows[1:]]


def _plain(v):
    return v.isoformat() if isinstance(v, (dt.date, dt.datetime, dt.time)) else v


def _records(data) -> list[dict]:
    """Coerce loaded data into a list of flat records for tabular output."""
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list)]
        data = lists[0] if len(data) == 1 and lists else [data]
    if not isinstance(data, list):
        raise ConversionError("Data must be a list of records to become a table")
    out = []
    for item in data:
        if isinstance(item, dict):
            out.append({k: _cell(v) for k, v in _flatten(item).items()})
        else:
            out.append({"value": _cell(item)})
    return out


def _flatten(d: dict, prefix: str = "") -> dict:
    flat = {}
    for k, v in d.items():
        key = f"{prefix}.{k}" if prefix else str(k)
        if isinstance(v, dict):
            flat.update(_flatten(v, key))
        else:
            flat[key] = v
    return flat


def _cell(v):
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False, default=str)
    return _plain(v)


def _drop_none(v):
    if isinstance(v, dict):
        return {str(k): _drop_none(x) for k, x in v.items() if x is not None}
    if isinstance(v, list):
        return [_drop_none(x) for x in v if x is not None]
    return v


def dump(data, path: Path, fmt: str) -> None:
    if fmt == "json":
        path.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    elif fmt == "jsonl":
        if isinstance(data, dict):
            lists = [v for v in data.values() if isinstance(v, list)]
            data = lists[0] if len(data) == 1 and lists else [data]
        items = data if isinstance(data, list) else [data]
        path.write_text("".join(json.dumps(i, ensure_ascii=False, default=str) + "\n" for i in items),
                        encoding="utf-8")
    elif fmt == "plist":
        with path.open("wb") as f:
            plistlib.dump(_drop_none(json.loads(json.dumps(data, default=str))), f)
    elif fmt == "yaml":
        path.write_text(yaml.safe_dump(json.loads(json.dumps(data, default=str)), sort_keys=False,
                                       allow_unicode=True), encoding="utf-8")
    elif fmt == "toml":
        if not isinstance(data, dict):
            data = {"items": data}
        path.write_text(tomli_w.dumps(_drop_none(data)), encoding="utf-8")
    elif fmt == "xml":
        if isinstance(data, dict) and len(data) == 1:
            (tag, val), = data.items()
        else:
            tag, val = "root", ({"item": data} if isinstance(data, list) else data)
        holder = ET.Element("_")
        _obj_to_xml(holder, tag, val)
        tree = ET.ElementTree(holder[0])
        ET.indent(tree)
        tree.write(path, encoding="utf-8", xml_declaration=True)
    elif fmt == "ini":
        if isinstance(data, list):  # table rows: one numbered section each
            data = {str(i): row for i, row in enumerate(_records(data), 1)}
        if not isinstance(data, dict):
            raise ConversionError("INI needs key/value sections")
        cp = configparser.ConfigParser(interpolation=None)
        cp.optionxform = str
        loose ={k: v for k, v in data.items() if not isinstance(v, dict)}
        if loose:
            cp["DEFAULT"] = {k: _scalar(_cell(v)) for k, v in loose.items()}
        for sec, vals in data.items():
            if isinstance(vals, dict):
                cp[str(sec)] = {k: _scalar(_cell(v)) for k, v in _flatten(vals).items()}
        with path.open("w", encoding="utf-8") as f:
            cp.write(f)
    elif fmt in ("csv", "tsv"):
        recs = _records(data)
        fields = list(dict.fromkeys(k for r in recs for k in r))
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields, delimiter="\t" if fmt == "tsv" else ",")
            w.writeheader()
            w.writerows(recs)
    elif fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        sheets = data if isinstance(data, dict) and data and all(isinstance(v, list) for v in data.values()) \
            and len(data) > 1 else {"Sheet1": data}
        wb.remove(wb.active)
        for name, rows in sheets.items():
            ws = wb.create_sheet(str(name)[:31])
            recs = _records(rows)
            fields = list(dict.fromkeys(k for r in recs for k in r))
            ws.append(fields)
            for r in recs:
                ws.append([r.get(k) for k in fields])
        wb.save(path)
    else:
        raise ConversionError(f"Can't write {fmt}")


@converter(inputs=TREE | TABLE | READ_ONLY, outputs=TREE | TABLE)
def data_convert(src: Path, dst: Path, src_fmt: str, dst_fmt: str) -> None:
    data = load(src, src_fmt)
    if src_fmt in ("csv", "tsv") and dst_fmt == "xlsx":
        data = _numbers(data)
    dump(data, dst, dst_fmt)
